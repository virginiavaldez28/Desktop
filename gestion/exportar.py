"""Exportación de reportes a Excel (.xlsx) y PDF."""
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .formato import moneda, numero, porcentaje
from .reportes import ENTERO, MONEDA, PORCENTAJE, TEXTO

EMPRESA = "Grupo Valpob S.R.L."


def texto_celda(tabla, fila, indice, valor):
    if valor is None:
        return ""
    formato = tabla.formato(fila, indice)
    if formato == PORCENTAJE:
        return porcentaje(valor)
    if formato == ENTERO:
        return numero(valor, 0)
    if formato == TEXTO or not isinstance(valor, (int, float, Decimal)):
        return str(valor)
    return moneda(valor)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def a_excel(reporte):
    libro = Workbook()
    hoja = libro.active
    hoja.title = reporte.titulo[:31]
    negrita = Font(bold=True)
    relleno_seccion = PatternFill("solid", fgColor="DCE6F1")
    relleno_total = PatternFill("solid", fgColor="F2F2F2")
    borde_arriba = Border(top=Side(style="thin"))

    hoja.append([f"{reporte.titulo.upper()} — {EMPRESA}"])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append([reporte.subtitulo])
    hoja.append([f"Generado el {timezone.localtime():%d/%m/%Y %H:%M}. Cifras en pesos argentinos."])
    for advertencia in reporte.advertencias:
        hoja.append([f"Atención: {advertencia}"])
    ancho_max = 1
    for tabla in reporte.tablas:
        hoja.append([])
        hoja.append([tabla.titulo])
        hoja.cell(hoja.max_row, 1).font = Font(bold=True, size=12)
        hoja.append([tabla.primera_columna] + list(tabla.columnas))
        fila_encabezado = hoja.max_row
        for celda in hoja[fila_encabezado]:
            celda.font = negrita
            celda.fill = relleno_seccion
            celda.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        for fila in tabla.filas:
            valores = []
            for indice, valor in enumerate(fila.valores):
                valores.append(float(valor) if isinstance(valor, Decimal) else valor)
            hoja.append([fila.etiqueta] + valores)
            n = hoja.max_row
            for indice in range(len(fila.valores)):
                celda = hoja.cell(n, indice + 2)
                formato = tabla.formato(fila, indice)
                if formato == PORCENTAJE:
                    celda.number_format = "0.0%"
                elif formato == ENTERO:
                    celda.number_format = "0"
                elif formato == MONEDA:
                    celda.number_format = '#,##0.00;[Red]-#,##0.00'
            if fila.estilo == "seccion":
                for celda in hoja[n]:
                    celda.font = negrita
                    celda.fill = relleno_seccion
            elif fila.estilo in ("subtotal", "total", "rubro"):
                for celda in hoja[n]:
                    celda.font = negrita
                    if fila.estilo == "total":
                        celda.fill = relleno_total
                        celda.border = borde_arriba
            elif fila.estilo == "detalle":
                hoja.cell(n, 1).alignment = Alignment(indent=1)
        if tabla.nota:
            hoja.append([tabla.nota])
            hoja.cell(hoja.max_row, 1).font = Font(italic=True, size=9)
        ancho_max = max(ancho_max, len(tabla.columnas) + 1)
    hoja.column_dimensions["A"].width = 55
    for columna in range(2, ancho_max + 1):
        hoja.column_dimensions[get_column_letter(columna)].width = 19
    hoja.freeze_panes = "B1"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def a_pdf(reporte):
    salida = BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=f"{reporte.titulo} — {reporte.subtitulo}", author=EMPRESA,
    )
    estilo_titulo = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=14, leading=17)
    estilo_sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#52514e"))
    estilo_tabla = ParagraphStyle("tabla", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=6)
    estilo_nota = ParagraphStyle("nota", fontName="Helvetica-Oblique", fontSize=7.5, leading=10)
    ancho_util = landscape(A4)[0] - 20 * mm

    historia = [
        Paragraph(escape(f"{reporte.titulo} — {EMPRESA}"), estilo_titulo),
        Paragraph(f"{escape(reporte.subtitulo)} · Generado el {timezone.localtime():%d/%m/%Y %H:%M} · Cifras en pesos argentinos", estilo_sub),
    ]
    for advertencia in reporte.advertencias:
        historia.append(Paragraph(escape(f"Atención: {advertencia}"), estilo_nota))
    historia.append(Spacer(1, 4 * mm))

    for tabla in reporte.tablas:
        columnas = len(tabla.columnas)
        tamanio = 8 if columnas <= 7 else 6.5
        estilo_celda = ParagraphStyle("celda", fontName="Helvetica", fontSize=tamanio, leading=tamanio + 2)
        estilo_enc = ParagraphStyle("enc", parent=estilo_celda, fontName="Helvetica-Bold", alignment=1)
        ancho_etiqueta = min(78 * mm, ancho_util * (0.34 if columnas > 7 else 0.4))
        ancho_columna = (ancho_util - ancho_etiqueta) / max(columnas, 1)
        if columnas == 1:
            ancho_etiqueta, ancho_columna = 110 * mm, 60 * mm

        datos = [[Paragraph(escape(tabla.primera_columna), estilo_enc)]
                 + [Paragraph(escape(c), estilo_enc) for c in tabla.columnas]]
        estilos = [
            ("FONT", (0, 0), (-1, -1), "Helvetica", tamanio),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#52514e")),
            ("TOPPADDING", (0, 0), (-1, -1), 1.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ]
        for fila in tabla.filas:
            n = len(datos)
            sangria = "&nbsp;&nbsp;&nbsp;" if fila.estilo == "detalle" else ""
            negrita = fila.estilo in ("seccion", "subtotal", "total", "rubro")
            etiqueta = f"<b>{escape(fila.etiqueta)}</b>" if negrita else sangria + escape(fila.etiqueta)
            datos.append([Paragraph(etiqueta, estilo_celda)] + [
                texto_celda(tabla, fila, i, v) for i, v in enumerate(fila.valores)
            ])
            if negrita:
                estilos.append(("FONT", (1, n), (-1, n), "Helvetica-Bold", tamanio))
            if fila.estilo == "seccion":
                estilos.append(("BACKGROUND", (0, n), (-1, n), colors.HexColor("#EEF3F9")))
            if fila.estilo == "total":
                estilos.append(("BACKGROUND", (0, n), (-1, n), colors.HexColor("#F2F2F2")))
                estilos.append(("LINEABOVE", (0, n), (-1, n), 0.5, colors.HexColor("#52514e")))
        tabla_pdf = Table(datos, colWidths=[ancho_etiqueta] + [ancho_columna] * columnas, repeatRows=1)
        tabla_pdf.setStyle(TableStyle(estilos))
        historia += [Paragraph(escape(tabla.titulo), estilo_tabla), Spacer(1, 2 * mm), tabla_pdf]
        if tabla.nota:
            historia += [Spacer(1, 2 * mm), Paragraph(escape(tabla.nota), estilo_nota)]
        historia.append(Spacer(1, 5 * mm))
    documento.build(historia)
    return salida.getvalue()
