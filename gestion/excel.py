"""Lectura del Excel original "Valpob_Punto_Equilibrio_Rentabilidad.xlsx".

Se usa para dos cosas:
  * la carga inicial (`python manage.py importar_excel archivo.xlsx`), y
  * la verificación contra el Excel (`python manage.py comparar_excel archivo.xlsx`).

Las columnas se buscan por su título (fila de encabezados), no por su letra,
así un cambio menor en el orden de las columnas del Excel no rompe la lectura.
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

import openpyxl

from .catalogos import (
    MES_POR_NOMBRE,
    AsignacionPersonal,
    CategoriaCompra,
    RubroGasto,
    Servicio,
    ServicioCompra,
)

SERVICIO_POR_NOMBRE = {s.label: s for s in Servicio}
CATEGORIA_POR_NOMBRE = {c.label: c for c in CategoriaCompra}
SERVICIO_COMPRA_POR_NOMBRE = {s.label: s for s in ServicioCompra}

# Secciones de "Apertura de Costos y Gastos" que se cargan a mano, ítem por ítem.
RUBRO_POR_SECCION_APERTURA = {
    "Alquileres": RubroGasto.ALQUILERES,
    "Servicios (luz, gas, internet, etc.)": RubroGasto.SERVICIOS,
    "Seguros": RubroGasto.SEGUROS,
    "Impuestos y Tasas": RubroGasto.IMPUESTOS_TASAS,
    "Otros Gastos Operativos": RubroGasto.OTROS_GASTOS_OPERATIVOS,
    "Leasing Vehículos/Equipos": RubroGasto.LEASING,
    "Cuotas Planes de Pago": RubroGasto.PLANES_PAGO,
}
# Secciones de Apertura que se calculan solas (no se importan).
SECCIONES_AUTOMATICAS_APERTURA = {
    "Mano de Obra Directa", "Materiales e Insumos", "Combustible", "Mantenimiento de Equipos",
    "Otros Costos Directos", "Sueldos y Cargas Sociales (Administración)",
}
# Renglones del Estado de Resultados que se cargan a mano.
RUBRO_POR_RENGLON_ER = {
    "Comisiones Bancarias": RubroGasto.COMISIONES_BANCARIAS,
    "Intereses por Financiación": RubroGasto.INTERESES,
    "Impuesto a los Créditos y Débitos (Ley 25.413)": RubroGasto.IMPUESTO_DEBITOS_CREDITOS,
    "Otros Ingresos": RubroGasto.OTROS_INGRESOS,
    "Otros Egresos": RubroGasto.OTROS_EGRESOS_NO_OPERATIVOS,
    "Impuesto a las Ganancias (estimado)": RubroGasto.IMPUESTO_GANANCIAS,
}
RENGLON_SALDO_PROVEEDORES = "Saldo de Proveedores al cierre del mes"


class ErrorExcel(Exception):
    pass


def decimal(valor):
    if valor is None or valor == "":
        return Decimal("0")
    # Se pasa por str para no arrastrar la representación binaria del float de Excel.
    return Decimal(str(valor)).quantize(Decimal("0.01"))


def texto(valor):
    return "" if valor is None else str(valor).strip()


def comentario(celda):
    return celda.comment.text.strip() if celda.comment else ""


def a_fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    raise ErrorExcel(f"Fecha inválida: {valor!r}")


def es_ejemplo(*textos):
    return any("ejemplo" in texto(t).lower() for t in textos)


@dataclass
class RenglonNomina:
    fila: int
    apellido_nombre: str
    cuil: str
    mes: int
    mano_obra_directa: bool
    asignacion: str
    porcentaje: Decimal
    regimen: str
    haberes: Decimal
    contribuciones: Decimal
    sindicato: Decimal
    honorarios: Decimal
    observaciones: str


@dataclass
class RenglonComprobante:
    fila: int
    fecha: date
    tercero: str
    tipo_comprobante: str
    punto_venta: int
    numero: int
    mes: int
    monto: Decimal
    observaciones: str
    servicio: str = ""
    categoria: str = ""
    servicio_asignado: str = ""
    ejemplo: bool = False


@dataclass
class RenglonGasto:
    mes: int
    rubro: str
    concepto: str
    monto: Decimal
    observaciones: str
    ejemplo: bool = False


@dataclass
class ContenidoExcel:
    nomina: list = field(default_factory=list)
    ventas: list = field(default_factory=list)
    compras: list = field(default_factory=list)
    gastos: list = field(default_factory=list)
    saldo_proveedores: dict = field(default_factory=dict)
    modo_asignacion_manual: bool = False

    @property
    def meses(self):
        todos = {r.mes for r in self.nomina} | {r.mes for r in self.ventas} | {r.mes for r in self.compras}
        todos |= {g.mes for g in self.gastos} | set(self.saldo_proveedores)
        return sorted(todos)


def _encabezados(hoja, fila):
    return {texto(c.value).replace("\n", " "): c.column for c in hoja[fila] if c.value}


def _columna(encabezados, *nombres):
    for nombre in nombres:
        if nombre in encabezados:
            return encabezados[nombre]
    raise ErrorExcel(f"No encontré la columna {nombres[0]!r}. Columnas encontradas: {list(encabezados)}")


def _buscar_fila_encabezado(hoja, primer_titulo):
    for fila in range(1, 15):
        if texto(hoja.cell(fila, 1).value) == primer_titulo:
            return fila
    raise ErrorExcel(f'No encontré el encabezado "{primer_titulo}" en la hoja "{hoja.title}".')


def _mes(valor, hoja, fila):
    nombre = texto(valor)
    if nombre not in MES_POR_NOMBRE:
        raise ErrorExcel(f'Hoja "{hoja.title}", fila {fila}: mes inválido {nombre!r}.')
    return MES_POR_NOMBRE[nombre]


def leer_nomina(libro):
    hoja = libro["Personal y Nómina"]
    fila_enc = _buscar_fila_encabezado(hoja, "Apellido y Nombre")
    enc = _encabezados(hoja, fila_enc)
    col = {
        "nombre": _columna(enc, "Apellido y Nombre"),
        "cuil": _columna(enc, "CUIL"),
        "mes": _columna(enc, "Mes"),
        "mod": _columna(enc, "¿Mano de Obra Directa?"),
        "servicio": _columna(enc, "Servicio Asignado"),
        "pct": _columna(enc, "% Afectación a ese Servicio"),
        "regimen": _columna(enc, "Régimen / Vínculo"),
        "haberes": _columna(enc, "Haberes (Neto de Recibo)"),
        "contrib": _columna(enc, "Contribuciones Patronales"),
        "sindicato": _columna(enc, "Sindicato y Mutual"),
        "honorarios": _columna(enc, "Honorarios (Monotributista)"),
        "obs": _columna(enc, "Observaciones"),
    }
    asignacion_por_nombre = {a.label: a for a in AsignacionPersonal}
    asignacion_por_nombre["Administración"] = AsignacionPersonal.ADMINISTRACION
    renglones = []
    for fila in range(fila_enc + 1, hoja.max_row + 1):
        nombre = texto(hoja.cell(fila, col["nombre"]).value)
        if not nombre:
            continue
        servicio = texto(hoja.cell(fila, col["servicio"]).value)
        if servicio not in asignacion_por_nombre:
            raise ErrorExcel(f"Personal y Nómina, fila {fila}: servicio asignado desconocido {servicio!r}.")
        observaciones = texto(hoja.cell(fila, col["obs"]).value)
        nota = comentario(hoja.cell(fila, col["nombre"]))
        if nota:
            observaciones = f"{observaciones}\n\nNota del Excel: {nota}".strip()
        renglones.append(RenglonNomina(
            fila=fila,
            apellido_nombre=nombre,
            cuil=texto(hoja.cell(fila, col["cuil"]).value),
            mes=_mes(hoja.cell(fila, col["mes"]).value, hoja, fila),
            mano_obra_directa=texto(hoja.cell(fila, col["mod"]).value) == "Sí",
            asignacion=asignacion_por_nombre[servicio].value,
            porcentaje=(Decimal(str(hoja.cell(fila, col["pct"]).value or 0)) * 100).quantize(Decimal("0.01")),
            regimen=texto(hoja.cell(fila, col["regimen"]).value),
            haberes=decimal(hoja.cell(fila, col["haberes"]).value),
            contribuciones=decimal(hoja.cell(fila, col["contrib"]).value),
            sindicato=decimal(hoja.cell(fila, col["sindicato"]).value),
            honorarios=decimal(hoja.cell(fila, col["honorarios"]).value),
            observaciones=observaciones,
        ))
    return renglones


def _leer_comprobantes(hoja, columna_tercero):
    fila_enc = _buscar_fila_encabezado(hoja, "Fecha")
    enc = _encabezados(hoja, fila_enc)
    col = {
        "fecha": _columna(enc, "Fecha"),
        "tercero": _columna(enc, columna_tercero),
        "tipo": _columna(enc, "Tipo de Comprobante"),
        "pv": _columna(enc, "Punto de Venta"),
        "numero": _columna(enc, "N° de Factura"),
        "mes": _columna(enc, "Mes"),
        "monto": _columna(enc, "Monto"),
        "obs": _columna(enc, "Observaciones"),
    }
    for fila in range(fila_enc + 1, hoja.max_row + 1):
        if not any(hoja.cell(fila, c).value not in (None, "") for c in col.values()):
            continue
        yield fila, enc, RenglonComprobante(
            fila=fila,
            fecha=a_fecha(hoja.cell(fila, col["fecha"]).value),
            tercero=texto(hoja.cell(fila, col["tercero"]).value),
            tipo_comprobante=texto(hoja.cell(fila, col["tipo"]).value),
            punto_venta=int(hoja.cell(fila, col["pv"]).value or 0),
            numero=int(hoja.cell(fila, col["numero"]).value or 0),
            mes=_mes(hoja.cell(fila, col["mes"]).value, hoja, fila),
            monto=decimal(hoja.cell(fila, col["monto"]).value),
            observaciones=texto(hoja.cell(fila, col["obs"]).value),
        )


def leer_ventas(libro):
    hoja = libro["Registro de Ventas"]
    renglones = []
    for fila, enc, renglon in _leer_comprobantes(hoja, "Cliente"):
        servicio = texto(hoja.cell(fila, _columna(enc, "Servicio")).value)
        if servicio not in SERVICIO_POR_NOMBRE:
            raise ErrorExcel(f"Registro de Ventas, fila {fila}: servicio desconocido {servicio!r}.")
        renglon.servicio = SERVICIO_POR_NOMBRE[servicio].value
        renglones.append(renglon)
    return renglones


def leer_compras(libro):
    hoja = libro["Registro de Compras"]
    renglones = []
    for fila, enc, renglon in _leer_comprobantes(hoja, "Proveedor"):
        categoria = texto(hoja.cell(fila, _columna(enc, "Categoría")).value)
        servicio = texto(hoja.cell(fila, _columna(enc, "Servicio Asignado")).value)
        if categoria not in CATEGORIA_POR_NOMBRE:
            raise ErrorExcel(f"Registro de Compras, fila {fila}: categoría desconocida {categoria!r}.")
        if servicio not in SERVICIO_COMPRA_POR_NOMBRE:
            raise ErrorExcel(f"Registro de Compras, fila {fila}: servicio desconocido {servicio!r}.")
        renglon.categoria = CATEGORIA_POR_NOMBRE[categoria].value
        renglon.servicio_asignado = SERVICIO_COMPRA_POR_NOMBRE[servicio].value
        renglon.ejemplo = es_ejemplo(renglon.tercero, renglon.observaciones)
        renglones.append(renglon)
    return renglones


def _columnas_meses(hoja, fila_enc):
    """{columna: número de mes} según los títulos de la fila de encabezados."""
    return {
        c.column: MES_POR_NOMBRE[texto(c.value)]
        for c in hoja[fila_enc]
        if texto(c.value) in MES_POR_NOMBRE
    }


def leer_gastos_apertura(libro):
    hoja = libro["Apertura de Costos y Gastos"]
    fila_enc = _buscar_fila_encabezado(hoja, "Concepto")
    meses = _columnas_meses(hoja, fila_enc)
    gastos = []
    seccion = None
    for fila in range(fila_enc + 1, hoja.max_row + 1):
        etiqueta = texto(hoja.cell(fila, 1).value)
        if not etiqueta:
            continue
        if etiqueta in RUBRO_POR_SECCION_APERTURA or etiqueta in SECCIONES_AUTOMATICAS_APERTURA or etiqueta.isupper():
            seccion = etiqueta
            continue
        if etiqueta.startswith(("Subtotal", "Total automático")) or seccion not in RUBRO_POR_SECCION_APERTURA:
            continue
        for columna, mes in meses.items():
            celda = hoja.cell(fila, columna)
            if isinstance(celda.value, str) and celda.value.startswith("="):
                continue  # fórmula: no es un dato cargado a mano
            monto = decimal(celda.value)
            if monto:
                gastos.append(RenglonGasto(
                    mes=mes,
                    rubro=RUBRO_POR_SECCION_APERTURA[seccion].value,
                    concepto=etiqueta,
                    monto=monto,
                    observaciones=comentario(celda),
                    ejemplo=es_ejemplo(etiqueta),
                ))
    return gastos


def leer_estado_resultados_manual(libro):
    """Gastos Bancarios, Otros Ingresos/Egresos, Ganancias y Saldo de Proveedores."""
    hoja = libro["Estado de Resultados"]
    fila_enc = _buscar_fila_encabezado(hoja, "Concepto")
    meses = _columnas_meses(hoja, fila_enc)
    gastos, saldos = [], {}
    for fila in range(fila_enc + 1, hoja.max_row + 1):
        etiqueta = texto(hoja.cell(fila, 1).value)
        if etiqueta not in RUBRO_POR_RENGLON_ER and etiqueta != RENGLON_SALDO_PROVEEDORES:
            continue
        for columna, mes in meses.items():
            celda = hoja.cell(fila, columna)
            if isinstance(celda.value, str) and celda.value.startswith("="):
                continue
            if celda.value in (None, ""):
                continue
            monto = decimal(celda.value)
            if etiqueta == RENGLON_SALDO_PROVEEDORES:
                saldos[mes] = monto
            elif monto:
                gastos.append(RenglonGasto(
                    mes=mes, rubro=RUBRO_POR_RENGLON_ER[etiqueta].value, concepto=etiqueta,
                    monto=monto, observaciones=comentario(celda),
                ))
    return gastos, saldos


def leer_excel(ruta) -> ContenidoExcel:
    libro = openpyxl.load_workbook(ruta)  # con fórmulas, para distinguir datos de cálculos
    contenido = ContenidoExcel(
        nomina=leer_nomina(libro),
        ventas=leer_ventas(libro),
        compras=leer_compras(libro),
        gastos=leer_gastos_apertura(libro),
    )
    gastos_er, contenido.saldo_proveedores = leer_estado_resultados_manual(libro)
    contenido.gastos += gastos_er
    modo = texto(libro["Costos por Servicio"]["J25"].value)
    contenido.modo_asignacion_manual = modo == "Manual"
    return contenido
