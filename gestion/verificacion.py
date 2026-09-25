"""Comparación celda por celda contra los valores calculados del Excel original.

Compara tres cosas:

1. Estado de Resultados, todos los meses con datos: tiene que dar igual.
2. Costos por Servicio, Punto de Equilibrio y Rentabilidad, para el mes elegido
   en el Excel (celda B2 de "Costos por Servicio"), calculados con el mismo
   criterio de mano de obra que usa la fórmula del Excel: tiene que dar igual.
   Esto prueba que el motor reproduce las fórmulas.
3. La única diferencia de criterio entre la aplicación y el Excel: la fórmula de
   Mano de Obra Directa del Excel (fila 10 de "Costos por Servicio") sólo suma
   personal asignado directo a Soporte PAE y a Cercos; si alguien está asignado
   directo a otro servicio (ej. Hidrolavadoras), el Excel lo cuenta en el Estado
   de Resultados pero no en Costos por Servicio. La aplicación lo suma a su
   servicio, así el Resultado Neto por Servicio cierra con el Estado de Resultados.
"""
import copy
from dataclasses import dataclass
from decimal import Decimal

import openpyxl

from .calculos import (
    RUBROS_GASTO_OPERATIVO,
    DatosMes,
    costos_por_servicio,
    estado_resultados,
    punto_equilibrio,
)
from .catalogos import MES_POR_NOMBRE, NOMBRE_MES, SERVICIOS, SERVICIOS_POOL, CategoriaCompra, RubroGasto
from .excel import SERVICIO_POR_NOMBRE, texto
from .models import Periodo

TOLERANCIA_MONTO = Decimal("0.01")
TOLERANCIA_PORCENTAJE = Decimal("0.0000001")


@dataclass
class Linea:
    hoja: str
    concepto: str
    columna: str
    excel: Decimal
    app: Decimal
    es_porcentaje: bool = False

    @property
    def diferencia(self):
        return self.app - self.excel

    @property
    def coincide(self):
        tolerancia = TOLERANCIA_PORCENTAJE if self.es_porcentaje else TOLERANCIA_MONTO
        return abs(self.diferencia) <= tolerancia


@dataclass
class ResultadoVerificacion:
    lineas: list
    mes_costos: str
    # {servicio: $} personal asignado directo a servicios del pool, que el Excel no suma en Costos por Servicio.
    mano_obra_directa_omitida: dict
    resultado_por_servicio_app: Decimal
    resultado_por_servicio_excel: Decimal
    resultado_operativo: Decimal

    @property
    def diferencias(self):
        return [linea for linea in self.lineas if not linea.coincide]

    @property
    def todo_coincide(self):
        return not self.diferencias


def _num(valor):
    return Decimal(str(valor or 0))


def _filas_por_etiqueta(hoja):
    return {texto(hoja.cell(f, 1).value): f for f in range(1, hoja.max_row + 1) if hoja.cell(f, 1).value}


def _valores_estado_resultados(er):
    valores = {s.label: er.ventas[s] for s in SERVICIOS}
    valores.update({
        "Total Ventas": er.total_ventas,
        "Mano de Obra Directa": er.mano_obra_directa,
        "Materiales e Insumos": er.materiales,
        "Combustible": er.combustible,
        "Mantenimiento de Equipos": er.mantenimiento,
        "Otros Costos Directos": er.otros_costos_directos,
        "Total Costo de Servicios Prestados": er.total_costo_servicios,
        "UTILIDAD BRUTA": er.utilidad_bruta,
        "Sueldos y Cargas Sociales (Administración)": er.sueldos_administracion,
        "Total Gastos Operativos": er.total_gastos_operativos,
        "Total Gastos Bancarios": er.total_gastos_bancarios,
        "RESULTADO OPERATIVO": er.resultado_operativo,
        "Otros Ingresos": er.otros_ingresos,
        "Otros Egresos": er.otros_egresos,
        "RESULTADO ANTES DE IMPUESTOS": er.resultado_antes_impuestos,
        "Impuesto a las Ganancias (estimado)": er.impuesto_ganancias,
        "RESULTADO NETO DEL PERÍODO": er.resultado_neto,
    })
    for rubro in RUBROS_GASTO_OPERATIVO:
        valores[rubro.label] = er.gastos_operativos[rubro]
    for rubro, monto in er.gastos_bancarios.items():
        valores[RubroGasto(rubro).label] = monto
    return valores


def _datos_criterio_excel(datos):
    """Copia de los datos del mes sin el personal asignado directo a servicios del pool."""
    copia = copy.deepcopy(datos)
    for s in SERVICIOS_POOL:
        copia.mano_obra[s.value] = Decimal("0")
    return copia


def _comparar_costos(hoja, cxs, nombre_hoja="Costos por Servicio"):
    lineas = []
    filas = _filas_por_etiqueta(hoja)
    columnas = {texto(c.value): c.column for c in hoja[4] if c.value}
    filas_cv = {
        "Mano de Obra Directa": "mano_obra",
        "Materiales e Insumos": CategoriaCompra.MATERIALES,
        "Combustible": CategoriaCompra.COMBUSTIBLE,
        "Mantenimiento de Equipos Específicos": CategoriaCompra.MANTENIMIENTO_EQUIPOS,
        "Otros Costos Variables": CategoriaCompra.OTROS_COSTOS_DIRECTOS,
    }
    por_servicio = {
        "Ventas Netas": (cxs.ventas, False),
        "Total Costos Variables": (cxs.total_costos_variables, False),
        "MARGEN DE CONTRIBUCIÓN ($)": (cxs.margen_contribucion, False),
        "MARGEN DE CONTRIBUCIÓN (%)": ({s: cxs.mc_porcentaje(s) for s in SERVICIOS}, True),
        "Costos Fijos Asignados ($)": (cxs.costos_fijos_asignados, False),
        "RESULTADO NETO POR SERVICIO": (cxs.resultado_neto, False),
        "RENTABILIDAD NETA (%)": ({s: cxs.rentabilidad(s) for s in SERVICIOS}, True),
    }
    for etiqueta, fila_cv in filas_cv.items():
        por_servicio[etiqueta] = (cxs.costos_variables[fila_cv], False)
    etiqueta_pct = next(e for e in filas if e.startswith("% de Asignación"))
    por_servicio[etiqueta_pct] = (cxs.porcentaje_asignacion, True)

    totales = {
        "Ventas Netas": (cxs.total_ventas, False),
        "Total Costos Variables": (cxs.total_cv, False),
        "MARGEN DE CONTRIBUCIÓN ($)": (cxs.total_mc, False),
        "MARGEN DE CONTRIBUCIÓN (%)": (cxs.mc_porcentaje(), True),
        "Costos Fijos Asignados ($)": (sum(cxs.costos_fijos_asignados.values()), False),
        "RESULTADO NETO POR SERVICIO": (cxs.total_resultado, False),
        "RENTABILIDAD NETA (%)": (cxs.rentabilidad(), True),
        "Total Gastos Operativos del Mes": (cxs.total_gastos_operativos, False),
        "Total Gastos Bancarios del Mes": (cxs.total_gastos_bancarios, False),
        "Total Costos Fijos del Período": (cxs.costos_fijos, False),
        etiqueta_pct: (sum(cxs.porcentaje_asignacion.values()), True),
    }
    for etiqueta, fila_cv in filas_cv.items():
        totales[etiqueta] = (cxs.total_fila(fila_cv), False)
    etiqueta_pool = next(e for e in filas if e.startswith("Mano de Obra Variable del Mes"))
    totales[etiqueta_pool] = (cxs.pool_operativo, False)

    for etiqueta, (valores, es_pct) in por_servicio.items():
        for s in SERVICIOS:
            celda = hoja.cell(filas[etiqueta], columnas[s.label])
            lineas.append(Linea(nombre_hoja, etiqueta, s.label, _num(celda.value), valores[s], es_pct))
    for etiqueta, (valor, es_pct) in totales.items():
        celda = hoja.cell(filas[etiqueta], columnas["TOTAL"])
        lineas.append(Linea(nombre_hoja, etiqueta, "TOTAL", _num(celda.value), valor, es_pct))
    return lineas


def _comparar_punto_equilibrio(hoja, cxs):
    general, por_servicio = punto_equilibrio(cxs)
    filas = _filas_por_etiqueta(hoja)
    lineas = []
    generales = {
        "Ventas Totales del Período": (general.ventas, False),
        "Costos Variables Totales": (general.costos_variables, False),
        "Margen de Contribución Total ($)": (general.margen_contribucion, False),
        "Margen de Contribución (%) Ponderado": (general.mc_porcentaje, True),
        "Costos Fijos Totales del Período": (general.costos_fijos, False),
        "PUNTO DE EQUILIBRIO EN PESOS ($)": (general.punto_equilibrio, False),
        "Diferencia vs. Punto de Equilibrio ($)": (general.diferencia, False),
        "Margen de Seguridad (%)": (general.margen_seguridad, True),
    }
    for etiqueta, (valor, es_pct) in generales.items():
        lineas.append(Linea("Punto de Equilibrio", etiqueta, "General", _num(hoja.cell(filas[etiqueta], 2).value), valor, es_pct))
    for s in SERVICIOS:
        fila = filas[s.label]
        pe = por_servicio[s]
        lineas.append(Linea("Punto de Equilibrio", "Punto de Equilibrio ($)", s.label, _num(hoja.cell(fila, 5).value), pe.punto_equilibrio))
        lineas.append(Linea("Punto de Equilibrio", "Margen de Seguridad ($)", s.label, _num(hoja.cell(fila, 6).value), pe.diferencia))
    return lineas


def verificar(ruta, anio=2026) -> ResultadoVerificacion:
    libro = openpyxl.load_workbook(ruta, data_only=True)
    lineas = []

    # 1. Estado de Resultados, todos los meses que existen en la aplicación.
    hoja_er = libro["Estado de Resultados"]
    filas_er = _filas_por_etiqueta(hoja_er)
    for celda in hoja_er[3]:
        mes = MES_POR_NOMBRE.get(texto(celda.value))
        periodo = mes and Periodo.objects.filter(anio=anio, mes=mes).first()
        if not periodo:
            continue
        valores = _valores_estado_resultados(estado_resultados(DatosMes.desde_base(periodo)))
        for etiqueta, valor in valores.items():
            excel = _num(hoja_er.cell(filas_er[etiqueta], celda.column).value)
            lineas.append(Linea("Estado de Resultados", etiqueta, str(periodo), excel, valor))

    # 2. Costos por Servicio / Punto de Equilibrio / Rentabilidad con el criterio del Excel.
    hoja_cxs = libro["Costos por Servicio"]
    nombre_mes = texto(hoja_cxs["B2"].value)
    periodo = Periodo.objects.get(anio=anio, mes=MES_POR_NOMBRE[nombre_mes])
    datos = DatosMes.desde_base(periodo)
    er = estado_resultados(datos)
    cxs_excel = costos_por_servicio(_datos_criterio_excel(datos), er)
    lineas += _comparar_costos(hoja_cxs, cxs_excel)
    lineas += _comparar_punto_equilibrio(libro["Punto de Equilibrio"], cxs_excel)

    hoja_rent = libro["Rentabilidad por Servicio"]
    filas_rent = _filas_por_etiqueta(hoja_rent)
    for s in SERVICIOS:
        fila = filas_rent[s.label]
        lineas.append(Linea("Rentabilidad por Servicio", "Resultado Neto", s.label, _num(hoja_rent.cell(fila, 7).value), cxs_excel.resultado_neto[s]))
        lineas.append(Linea("Rentabilidad por Servicio", "Rentabilidad Neta (%)", s.label, _num(hoja_rent.cell(fila, 8).value), cxs_excel.rentabilidad(s), True))

    # 3. La diferencia de criterio.
    cxs_app = costos_por_servicio(datos, er)
    omitida = {s: datos.mano_obra[s.value] for s in SERVICIOS_POOL if datos.mano_obra[s.value]}
    return ResultadoVerificacion(
        lineas=lineas,
        mes_costos=f"{NOMBRE_MES[periodo.mes]} {periodo.anio}",
        mano_obra_directa_omitida=omitida,
        resultado_por_servicio_app=cxs_app.total_resultado,
        resultado_por_servicio_excel=_num(hoja_cxs.cell(_filas_por_etiqueta(hoja_cxs)["RESULTADO NETO POR SERVICIO"], 11).value),
        resultado_operativo=er.resultado_operativo,
    )
